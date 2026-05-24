"""Generates data/templates/hazmat_chit.docx from the source form.

Source of truth: ``BLANK NEW HAZMAT ISSUE CHIT 2.0.docx`` at the repo
root. This script clones it and injects docxtpl placeholders into the
appropriate cells, then writes the result to ``settings.template_path``.

Re-run any time the source form changes; commit both the source .docx
and the regenerated template.

    PYTHONPATH=. .venv/bin/python scripts/prepare_template.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import _Cell

from app.config import REPO_ROOT, settings

SOURCE_DOCX = REPO_ROOT / "BLANK NEW HAZMAT ISSUE CHIT 2.0.docx"

# (row, col) → placeholder text. Rows / cols refer to the source form's
# single 25×4 table. Cells listed here are the empty value-cells next to
# bold field labels in the header block.
HEADER_PLACEHOLDERS = {
    (2, 1): "{{ location }}",
    (4, 1): "{{ name }}",
    (5, 1): "{{ datetime }}",
    (6, 1): "{{ workcenter }}",
    (7, 1): "{{ lpo }}",
}

# Line-items block on the source form: row 9 is the header
# (SPMIG/SPIN | NONMENCLATURE | NIIN | QTY) and rows 10–15 are six
# blank data rows. We turn rows 10–12 into a docxtpl row loop and
# delete rows 13–15; the renderer pads ``lines`` so the printed form
# still shows at least 6 rows.
ITEMS_HEADER_ROW = 9
ITEMS_FIRST_ROW = 10
ITEMS_LAST_ROW = 15


def _set_cell_text(cell: _Cell, text: str, bold: bool | None = None) -> None:
    """Replace the cell's first paragraph with a single run of ``text``."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    if bold is not None:
        run.bold = bold


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(
            f"Source form missing at {SOURCE_DOCX}. "
            "Place the original chit there before regenerating the template."
        )

    doc = Document(str(SOURCE_DOCX))
    table = doc.tables[0]

    for (row, col), placeholder in HEADER_PLACEHOLDERS.items():
        _set_cell_text(table.rows[row].cells[col], placeholder)

    # docxtpl row loop. The {%tr ... %} marker consumes its entire row,
    # so we use row 10 for the opening tag, row 11 as the repeating data
    # row, and row 12 for the closing tag.
    _set_cell_text(table.rows[ITEMS_FIRST_ROW].cells[0], "{%tr for line in lines %}")
    data_row = table.rows[ITEMS_FIRST_ROW + 1]
    _set_cell_text(data_row.cells[0], "{{ line.spmig }}")
    _set_cell_text(data_row.cells[1], "{{ line.nomenclature }}")
    _set_cell_text(data_row.cells[2], "{{ line.niin }}")
    _set_cell_text(data_row.cells[3], "{{ line.qty }}")
    _set_cell_text(table.rows[ITEMS_FIRST_ROW + 2].cells[0], "{%tr endfor %}")

    tbl = table._tbl
    for idx in range(ITEMS_LAST_ROW, ITEMS_FIRST_ROW + 2, -1):
        tbl.remove(table.rows[idx]._tr)

    out = settings.template_path
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
